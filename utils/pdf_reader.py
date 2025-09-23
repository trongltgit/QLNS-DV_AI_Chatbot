# utils/pdf_reader.py
from pypdf import PdfReader
import docx
import openpyxl

def extract_pdf_text(file_path: str) -> str:
    """Đọc toàn bộ text từ PDF"""
    reader = PdfReader(file_path)
    text = ""
    for page in reader.pages:
        if page.extract_text():
            text += page.extract_text() + "\n"
    return text


def extract_docx_text(file_path: str) -> str:
    """Đọc toàn bộ text từ DOCX"""
    doc = docx.Document(file_path)
    text = ""
    for para in doc.paragraphs:
        if para.text.strip():
            text += para.text.strip() + "\n"
    return text


def extract_excel_text(file_path: str) -> str:
    """Đọc toàn bộ text từ Excel (XLSX)"""
    wb = openpyxl.load_workbook(file_path, data_only=True)
    text = ""
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            line = " ".join([str(cell) for cell in row if cell is not None])
            if line.strip():
                text += line.strip() + "\n"
    return text


def extract_text(file_path: str) -> str:
    """Tự động chọn parser theo loại file"""
    file_path = file_path.lower()
    if file_path.endswith(".pdf"):
        return extract_pdf_text(file_path)
    elif file_path.endswith(".docx"):
        return extract_docx_text(file_path)
    elif file_path.endswith(".xlsx"):
        return extract_excel_text(file_path)
    else:
        raise ValueError("Định dạng file không được hỗ trợ. Chỉ hỗ trợ PDF, DOCX, XLSX.")


def chunk_text(text: str, chunk_size: int = 500) -> list[str]:
    """Chia nhỏ text thành các đoạn (chunk) theo số lượng từ"""
    words = text.split()
    chunks = []
    for i in range(0, len(words), chunk_size):
        chunk = " ".join(words[i:i+chunk_size])
        chunks.append(chunk)
    return chunks


def extract_chunks(file_path: str, chunk_size: int = 500) -> list[str]:
    """Pipeline: đọc file → trích text → chia thành chunks"""
    raw_text = extract_text(file_path)
    return chunk_text(raw_text, chunk_size)
